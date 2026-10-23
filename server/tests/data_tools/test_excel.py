import hashlib
import io
import json
import zipfile
from datetime import datetime
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from flowtrace_server.exports.service import xlsx_bytes
from flowtrace_server.imports.parsers import ImportFormatError, parse_xlsx
from flowtrace_server.imports.qianji_template import TemplateMapping, parse_qianji_xlsx


def document():
    return {
        "format": "flowtrace-export-v1",
        "document": {
            "format": "flowtrace-data-v1",
            "source_namespace": "synthetic",
            "rows": [
                {
                    "source_id": str(uuid4()),
                    "entity_type": "bill",
                    "scope_id": str(uuid4()),
                    "payload": {
                        "bill_type": "expense",
                        "money": {
                            "currency": "CNY",
                            "currency_scale": 2,
                            "amount_minor": "9007199254740993",
                        },
                        "business_date": "2026-10-21",
                        "business_timezone": "Asia/Shanghai",
                        "note": '=HYPERLINK("https://invalid.example","private-secret")',
                    },
                }
            ],
        },
        "attachments": [],
    }


def test_export_excel_and_import_preserve_exact_large_string_and_note():
    source = document()
    content = xlsx_bytes(source)
    parsed = parse_xlsx(content, "synthetic")
    assert not parsed.errors
    assert parsed.document.rows[0].payload == source["document"]["rows"][0]["payload"]
    book = load_workbook(io.BytesIO(content), data_only=False)
    try:
        assert all(cell.data_type != "f" for row in book["records"] for cell in row)
        assert json.loads(book["records"].cell(2, 4).value)["money"]["amount_minor"] == (
            "9007199254740993"
        )
    finally:
        book.close()


def test_excel_retains_signed_historical_posting_without_changing_payload():
    source = document()
    row = source["document"]["rows"][0]
    row["payload"].update(
        bill_type="balance_adjustment", account_id=str(uuid4()), target_balance_minor="10000"
    )
    hint = {"signed_principal_minor": "-9007199254740993"}
    source["document"]["posting_hints"] = {row["source_id"]: hint}
    parsed = parse_xlsx(xlsx_bytes(source), "synthetic")
    assert not parsed.errors
    assert parsed.document.rows[0].payload == row["payload"]
    assert parsed.document.posting_hints[row["source_id"]].model_dump() == hint


def test_excel_formula_cells_rejected_without_evaluation():
    book = Workbook()
    sheet = book.active
    sheet.title = "records"
    sheet.append(["source_id", "entity_type", "scope_id", "payload"])
    sheet.append([str(uuid4()), "bill", str(uuid4()), '=WEBSERVICE("https://invalid.example")'])
    output = io.BytesIO()
    book.save(output)
    result = parse_xlsx(output.getvalue(), "synthetic")
    assert result.document is None and result.errors[0]["row"] == 2


def test_excel_embeds_verified_image_and_keeps_authorized_link_manifest():
    source = document()
    image = io.BytesIO()
    Image.new("RGB", (3, 2), "green").save(image, format="PNG")
    original = image.getvalue()
    identifier = str(uuid4())
    source["attachments"] = [
        {
            "attachment_id": identifier,
            "sha256": hashlib.sha256(original).hexdigest(),
            "size_bytes": len(original),
            "media_type": "image/png",
            "href": f"/api/v1/attachments/{identifier}/content",
        }
    ]
    content = xlsx_bytes(source, image_loader=lambda _: original, embed_images=True)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = archive.namelist()
        assert len([n for n in names if n.startswith("xl/media/")]) == 1
        assert archive.read(next(n for n in names if n.startswith("xl/media/"))) == original
    book = load_workbook(io.BytesIO(content))
    try:
        assert book["attachments"].cell(2, 2).value == hashlib.sha256(original).hexdigest()
    finally:
        book.close()


def test_excel_image_mismatch_fails_before_result():
    from flowtrace_server.auth.errors import ApiError

    source = document()
    source["attachments"] = [
        {
            "attachment_id": str(uuid4()),
            "sha256": "0" * 64,
            "size_bytes": 1,
            "media_type": "image/png",
            "href": "/safe",
        }
    ]
    with pytest.raises(ApiError, match="export_image_hash_mismatch"):
        xlsx_bytes(source, image_loader=lambda _: b"x", embed_images=True)


def test_public_qianji_template_excel_datetime_and_exact_numeric():
    book = Workbook()
    sheet = book.active
    sheet.append(["时间", "类型", "金额", "备注"])
    sheet.append([datetime(2026, 10, 21, 12, 15), "收入", 1.25, "合成Excel"])
    stream = io.BytesIO()
    book.save(stream)
    parsed = parse_qianji_xlsx(
        stream.getvalue(), TemplateMapping(scope_id=str(uuid4()), timezone="Asia/Shanghai")
    )
    assert not parsed.errors
    payload = parsed.document.rows[0].payload
    assert payload["money"]["amount_minor"] == "125"
    assert payload["occurred_at"] == "2026-10-21T12:15:00+08:00"


def test_workbook_zip_expansion_limit_before_xml_parse():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("huge.xml", b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(ImportFormatError, match="workbook_expansion_too_large"):
        parse_xlsx(output.getvalue(), "synthetic")
