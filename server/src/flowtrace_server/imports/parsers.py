"""Bounded adapters for our documented format. No guessed Qianji compatibility."""

import csv
import hashlib
import io
import json
from dataclasses import dataclass

from pydantic import ValidationError

from .schemas import ImportDocument, ImportRow

MAX_FILE_BYTES = 5 * 1024 * 1024
HEADERS = ("source_id", "entity_type", "scope_id", "payload")


class ImportFormatError(ValueError):
    def __init__(self, code: str, row: int | None = None, field: str | None = None):
        self.code, self.row, self.field = code, row, field
        super().__init__(code)


@dataclass(frozen=True)
class ParsedTable:
    document: ImportDocument | None
    errors: tuple[dict, ...]


def canonical(value) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    ).encode("utf-8")


def digest(value) -> str:
    return hashlib.sha256(canonical(value)).hexdigest()


def _unique(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise ImportFormatError("duplicate_json_key", field=key)
        value[key] = item
    return value


def decode_json(content: bytes):
    return _decode_json(content, MAX_FILE_BYTES)


def _decode_json(content: bytes, max_file_bytes: int):
    """Shared decoder; only trusted stage code supplies the separate budget."""
    if len(content) > max_file_bytes:
        raise ImportFormatError("file_too_large")
    try:
        value = json.loads(
            content.decode("utf-8-sig"),
            object_pairs_hook=_unique,
            parse_float=lambda _: (_ for _ in ()).throw(ValueError()),
            parse_constant=lambda _: (_ for _ in ()).throw(ValueError()),
        )
        _validate_json(value)
        return value
    except (UnicodeError, ValueError, RecursionError) as error:
        if isinstance(error, ImportFormatError):
            raise
        raise ImportFormatError("invalid_json") from None


def _validate_json(value, depth=0):
    if depth > 32:
        raise ImportFormatError("json_too_deep")
    if isinstance(value, str):
        if "\x00" in value:
            raise ImportFormatError("invalid_json_text")
        value.encode("utf-8", errors="strict")
    elif isinstance(value, dict):
        for key, item in value.items():
            _validate_json(key, depth + 1)
            _validate_json(item, depth + 1)
    elif isinstance(value, list):
        for item in value:
            _validate_json(item, depth + 1)


def parse_json(content: bytes) -> ImportDocument:
    value = decode_json(content)
    if isinstance(value, dict) and value.get("format") == "flowtrace-export-v1":
        # An export is an envelope: attachments remain an explicit mapping step,
        # never network requests performed implicitly by the parser.
        allowed = {"format", "document", "attachments", "scope_versions", "filter"}
        if set(value) - allowed or "document" not in value:
            raise ImportFormatError("invalid_export_envelope")
        value = value["document"]
    return ImportDocument.model_validate(value)


def _table(rows, namespace: str) -> ParsedTable:
    accepted, errors = [], []
    for line, row in enumerate(rows, 2):
        if line > 51:
            errors.append({"row": line, "field": None, "code": "too_many_rows"})
            break
        try:
            if set(row) != set(HEADERS) or None in row.values():
                raise ImportFormatError("invalid_columns")
            candidate = dict(row)
            if not isinstance(candidate["payload"], str):
                raise ImportFormatError("payload_json_required", field="payload")
            candidate["payload"] = decode_json(candidate["payload"].encode())
            accepted.append(ImportRow.model_validate(candidate))
        except ImportFormatError as error:
            errors.append({"row": line, "field": error.field, "code": error.code})
        except ValidationError as error:
            for detail in error.errors(include_input=False, include_url=False):
                errors.append(
                    {
                        "row": line,
                        "field": ".".join(map(str, detail["loc"])),
                        "code": detail["type"],
                    }
                )
    if errors:
        return ParsedTable(None, tuple(errors))
    try:
        document = ImportDocument(
            format="flowtrace-data-v1", source_namespace=namespace, rows=accepted
        )
        return ParsedTable(document, ())
    except ValidationError as error:
        return ParsedTable(
            None,
            tuple(
                {"row": None, "field": ".".join(map(str, e["loc"])), "code": e["type"]}
                for e in error.errors(include_input=False)
            ),
        )


def parse_csv(content: bytes, namespace: str) -> ParsedTable:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFormatError("file_too_large")
    try:
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        if reader.fieldnames is None or tuple(reader.fieldnames) != HEADERS:
            raise ImportFormatError("invalid_columns", 1)
        return _table(reader, namespace)
    except (UnicodeError, csv.Error):
        raise ImportFormatError("invalid_csv") from None


def secure_workbook(content: bytes):
    return _secure_workbook(content, MAX_FILE_BYTES, 20 * 1024 * 1024, 200)


def _secure_workbook(content, max_file_bytes, max_expanded_bytes, max_archive_entries):
    """Same XML/ZIP protections, with a trusted caller's bounded budget."""
    if len(content) > max_file_bytes:
        raise ImportFormatError("file_too_large")
    # ZIP expansion is capped before the XML library receives any member.
    import zipfile

    from openpyxl import load_workbook
    from openpyxl.xml import DEFUSEDXML

    if not DEFUSEDXML:
        raise ImportFormatError("secure_xml_parser_required")

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > max_archive_entries
                or sum(m.file_size for m in members) > max_expanded_bytes
            ):
                raise ImportFormatError("workbook_expansion_too_large")
            if any(m.flag_bits & 1 or m.filename.endswith("vbaProject.bin") for m in members):
                raise ImportFormatError("workbook_active_content")
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ImportFormatError:
        raise
    except Exception:
        raise ImportFormatError("invalid_xlsx") from None


def parse_xlsx(content: bytes, namespace: str) -> ParsedTable:
    try:
        book = secure_workbook(content)
        try:
            if "records" not in book.sheetnames or len(book.sheetnames) > 5:
                raise ImportFormatError("records_sheet_required")
            sheet = book["records"]
            if sheet.max_column != 4 or sheet.max_row > 51:
                raise ImportFormatError("workbook_dimensions_invalid")
            rows = sheet.iter_rows()
            header = next(rows, ())
            if tuple(cell.value for cell in header) != HEADERS:
                raise ImportFormatError("invalid_columns", 1)

            def records():
                for cells in rows:
                    if any(cell.data_type == "f" for cell in cells):
                        yield {"payload": None}
                    else:
                        yield dict(zip(HEADERS, (cell.value for cell in cells), strict=True))

            parsed = _table(records(), namespace)
            if parsed.document is None or "historical_postings" not in book.sheetnames:
                return parsed
            postings = book["historical_postings"]
            if postings.max_column != 2 or postings.max_row > 51:
                raise ImportFormatError("historical_postings_dimensions_invalid")
            hint_rows = postings.iter_rows()
            if tuple(cell.value for cell in next(hint_rows, ())) != (
                "source_id",
                "signed_principal_minor",
            ):
                raise ImportFormatError("historical_postings_columns_invalid")
            hints = {}
            for cells in hint_rows:
                if any(cell.data_type == "f" or not isinstance(cell.value, str) for cell in cells):
                    raise ImportFormatError("historical_postings_values_invalid")
                identifier, signed = (cell.value for cell in cells)
                if identifier in hints:
                    raise ImportFormatError("duplicate_historical_posting")
                hints[identifier] = {"signed_principal_minor": signed}
            try:
                document = ImportDocument.model_validate(
                    {**parsed.document.model_dump(), "posting_hints": hints}
                )
            except ValidationError:
                raise ImportFormatError("invalid_historical_postings") from None
            return ParsedTable(document, ())
        finally:
            book.close()
    except ImportFormatError:
        raise
    except Exception:
        raise ImportFormatError("invalid_xlsx") from None
