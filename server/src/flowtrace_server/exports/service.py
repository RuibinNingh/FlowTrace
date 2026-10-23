import io
from collections.abc import Callable
from copy import deepcopy
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.orm import Session

from flowtrace_server.attachments.models import Attachment
from flowtrace_server.auth.dependencies import Principal
from flowtrace_server.auth.errors import ApiError
from flowtrace_server.auth.service import assert_active_principal
from flowtrace_server.books.permissions import authorize_scope
from flowtrace_server.imports.parsers import HEADERS, canonical
from flowtrace_server.imports.service import REFERENCE_FIELDS, REFERENCE_LISTS
from flowtrace_server.sync.common import lock_scopes, run_transaction
from flowtrace_server.sync.models import LedgerEntry, SyncEntity

from .schemas import ExportRequest


def _references(value, key=None):
    if isinstance(value, dict):
        return set().union(*(_references(v, k) for k, v in value.items()))
    if isinstance(value, list):
        return set().union(*(_references(v, key) for v in value))
    return {UUID(value)} if key in REFERENCE_FIELDS | REFERENCE_LISTS and value else set()


def export_document(session: Session, principal: Principal, body: ExportRequest) -> dict:
    """A consistent read requiring access to every complete financial reference.

    Shared-book readers lacking a private payer scope get a safe rejection. The
    raw financial payload is never partially reconstructed from public projections.
    """

    def read():
        if principal.purpose not in {"api_native", "api_web"}:
            raise ApiError("session_purpose_invalid", 403)
        assert_active_principal(session, principal, lock=True)
        requested = {UUID(i) for i in body.scope_ids}
        for identifier in requested:
            authorize_scope(session, principal, identifier)
        entities = list(
            session.scalars(
                select(SyncEntity)
                .where(SyncEntity.scope_id.in_(requested), SyncEntity.deleted.is_(False))
                .order_by(SyncEntity.id)
                .limit(1001)
            )
        )
        if len(entities) > 1000:
            raise ApiError("export_too_large", 413)
        selected = []
        for entity in entities:
            if entity.entity_type != "bill":
                if body.include_configuration:
                    selected.append(entity)
                continue
            day = entity.payload["business_date"]
            if (body.date_from and day < body.date_from) or (body.date_to and day > body.date_to):
                continue
            selected.append(entity)
        # Raw payloads include private payer/source IDs. Authorize all before copying anything.
        attachments = {}
        wanted = set(requested)
        for entity in selected:
            for identifier in _references(entity.payload):
                target = session.get(SyncEntity, identifier)
                if target is not None:
                    wanted.add(target.scope_id)
                else:
                    attachment = session.get(Attachment, identifier)
                    if attachment is None or attachment.status != "ready":
                        raise ApiError("export_reference_unavailable", 409)
                    wanted.add(attachment.scope_id)
                    attachments[str(identifier)] = {
                        "attachment_id": str(identifier),
                        "sha256": attachment.sha256,
                        "size_bytes": attachment.size_bytes,
                        "media_type": attachment.media_type,
                        "href": f"/api/v1/attachments/{identifier}/content",
                        "authorization": "current_session_and_scope_required",
                    }
        scopes = lock_scopes(session, wanted)
        for identifier in scopes:
            authorize_scope(session, principal, identifier)
        bills = [e for e in selected if e.entity_type == "bill"]
        bills.sort(
            key=lambda e: (e.payload["business_date"], str(e.id)), reverse=body.order == "date_desc"
        )
        selected = [e for e in selected if e.entity_type != "bill"] + bills
        rows = [
            {
                "source_id": str(e.id),
                "entity_type": e.entity_type,
                "scope_id": str(e.scope_id),
                "payload": deepcopy(e.payload),
            }
            for e in selected
        ]
        for row in rows:
            row["payload"].pop("preserve_hidden_financial_links", None)
        posting_hints = {}
        for entity in bills:
            if entity.payload["bill_type"] != "balance_adjustment":
                continue
            entries = list(
                session.scalars(
                    select(LedgerEntry).where(
                        LedgerEntry.bill_id == entity.id, LedgerEntry.component == "principal"
                    )
                )
            )
            if (
                len(entries) != 1
                or str(entries[0].account_id) != entity.payload["account_id"]
                or abs(entries[0].amount_minor) != int(entity.payload["money"]["amount_minor"])
                or entries[0].currency != entity.payload["money"]["currency"]
                or entries[0].currency_scale != entity.payload["money"]["currency_scale"]
            ):
                raise ApiError("export_historical_posting_inconsistent", 409)
            posting_hints[str(entity.id)] = {"signed_principal_minor": str(entries[0].amount_minor)}
        assert_active_principal(session, principal, lock=True)
        return {
            "format": "flowtrace-export-v1",
            "document": {
                "format": "flowtrace-data-v1",
                "source_namespace": "flowtrace-export",
                "rows": rows,
                "posting_hints": posting_hints,
            },
            "attachments": list(attachments.values()),
            "scope_versions": {
                str(k): {
                    "revision": str(v.revision),
                    "permission_version": str(v.permission_version),
                }
                for k, v in scopes.items()
            },
            "filter": body.model_dump(),
        }

    return run_transaction(session, read, repeatable=True)


def json_bytes(document: dict) -> bytes:
    return canonical(document)


def xlsx_bytes(
    document: dict,
    *,
    image_loader: Callable[[dict], bytes] | None = None,
    embed_images: bool = False,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image

    book = Workbook()
    records = book.active
    records.title = "records"
    records.append(HEADERS)
    for row in document["document"]["rows"]:
        payload = canonical(row["payload"]).decode()
        if len(payload) > 32767:
            raise ApiError("export_cell_too_large", 413)
        records.append([row["source_id"], row["entity_type"], row["scope_id"], payload])
        # Explicit string cells prevent notes from becoming formulas.
        for cell in records[records.max_row]:
            cell.data_type = "s"
    hints = document["document"].get("posting_hints", {})
    if hints:
        postings = book.create_sheet("historical_postings")
        postings.append(["source_id", "signed_principal_minor"])
        for identifier, hint in hints.items():
            postings.append([identifier, hint["signed_principal_minor"]])
            for cell in postings[postings.max_row]:
                cell.data_type = "s"
    pictures = book.create_sheet("attachments")
    pictures.append(["attachment_id", "sha256", "size_bytes", "media_type", "href"])
    total = 0
    for position, item in enumerate(document["attachments"], 2):
        pictures.append(
            [item[k] for k in ("attachment_id", "sha256", "size_bytes", "media_type", "href")]
        )
        if embed_images:
            # The authorized caller supplies verified original bytes, never a client URL/path.
            if image_loader is None:
                raise ApiError("export_image_store_required", 503)
            original = image_loader(item)
            if original is None:
                raise ApiError("export_image_store_required", 503)
            import hashlib

            if (
                len(original) != item["size_bytes"]
                or hashlib.sha256(original).hexdigest() != item["sha256"]
            ):
                raise ApiError("export_image_hash_mismatch", 409)
            total += len(original)
            if total > 20 * 1024 * 1024:
                raise ApiError("export_too_large", 413)
            picture = Image(io.BytesIO(original))
            factor = min(1, 200 / picture.width, 150 / picture.height)
            picture.width, picture.height = picture.width * factor, picture.height * factor
            pictures.add_image(picture, f"F{position}")
    output = io.BytesIO()
    book.save(output)
    book.close()
    if output.tell() > 25 * 1024 * 1024:
        raise ApiError("export_too_large", 413)
    return output.getvalue()
